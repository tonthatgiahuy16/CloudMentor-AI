from hashlib import sha256
from pathlib import Path
from uuid import uuid4

import numpy as np
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.db_models  # Đăng ký các bảng với Base.metadata
from app.api.subjects import get_subject_repository, router as subjects_router
from app.core.database import Base
from app.repos.subject_repository import SubjectRepository

from app.core import config
from app.rag.vector_store import VectorStore


from app.api.upload import (
    get_pdf_service,
    get_upload_dir,
    get_upload_subject_repository,
    get_source_storage,
    router as upload_router,
)
from app.models.chunk import Chunk
from app.repos.document_repository import DocumentRepository

from app.api.chat import (
    get_chat_service,
    get_chat_subject_repository,
    router as chat_router,
)
from app.rag.retriever import Retriever
from app.services.chat_service import ChatService

from app.api.documents import (
    get_document_deletion_service,
    get_document_repository,
    router as documents_router,
)
from app.services.document_deletion_service import (
    DocumentDeletionService,
)

from io import BytesIO

from docx import Document as WordDocument

from app.rag.loader import (
    read_docx_paragraphs,
)

from app.storage.source_storage import (
    LocalSourceStorage,
)


class LifecyclePDFService:
    """
    Thay thế các bước xử lý nặng trong test nhưng vẫn ghi vào
    database thật và Chroma thật đã được cô lập.
    """

    def __init__(
        self,
        document_repo,
        vector_store,
        source_storage,
    ):
        self.document_repo = document_repo
        self.vector_store = vector_store
        self.source_storage = source_storage

    def find_active_duplicate(
        self,
        *,
        subject_id,
        content_sha256,
    ):
        return (
            self.document_repo
            .get_active_by_subject_and_sha256(
                subject_id=subject_id,
                content_sha256=content_sha256,
            )
        )

    def upload(
        self,
        *,
        source_data,
        subject_id,
        chapter,
        document_id,
        original_filename,
        content_sha256,
        file_size_bytes,
    ):
        file_type = (
            Path(original_filename)
            .suffix
            .lstrip(".")
            .lower()
        )

        self.document_repo.create(
            document_id=document_id,
            subject_id=subject_id,
            filename=original_filename,
            file_type=file_type,
            chapter=chapter,
            source_uri=None,
            content_sha256=content_sha256,
            file_size_bytes=file_size_bytes,
        )

        stored_source = self.source_storage.store(
            document_id=document_id,
            filename=original_filename,
            data=source_data,
        )

        self.document_repo.attach_source_metadata(
            document_id=document_id,
            source_uri=stored_source.source_uri,
            content_sha256=stored_source.content_sha256,
            file_size_bytes=stored_source.file_size_bytes,
        )

        self.document_repo.mark_processing(document_id)

        if file_type == "txt":
            lines = source_data.decode(
                encoding="utf-8"
            ).splitlines()

            chunk_text = "\n".join(lines)
            source_location = {
                "line_range": (1, len(lines)),
            }

        elif file_type == "docx":
            paragraphs = read_docx_paragraphs(
                BytesIO(source_data)
            )

            chunk_text = "\n".join(paragraphs)
            source_location = {
                "paragraph_range": (
                    1,
                    len(paragraphs),
                ),
            }

        else:
            chunk_text = (
                "Cloud computing provides computing "
                "resources on demand."
            )
            source_location = {
                "page": 1,
            }

        chunk = Chunk(
            id=f"{document_id}:0",
            document_id=document_id,
            subject_id=subject_id,
            chunk_index=0,
            source=original_filename,
            text=chunk_text,
            **source_location,
        )

        embeddings = np.array(
            [[1.0, 0.0, 0.0]],
            dtype=np.float32,
        )

        self.vector_store.add(
            [chunk],
            embeddings,
        )
        self.document_repo.mark_indexed(document_id)

        return 1

    def retry_failed_document(
        self,
        document,
    ):
        source_data = self.source_storage.read_bytes(
            document.source_uri
        )

        self.vector_store.delete_by_document_id(
            document.document_id
        )
        self.document_repo.mark_processing(
            document.document_id
        )

        file_type = document.file_type

        if file_type == "txt":
            lines = source_data.decode(
                encoding="utf-8"
            ).splitlines()

            chunk_text = "\n".join(lines)
            source_location = {
                "line_range": (1, len(lines)),
            }

        elif file_type == "docx":
            paragraphs = read_docx_paragraphs(
                BytesIO(source_data)
            )

            chunk_text = "\n".join(paragraphs)
            source_location = {
                "paragraph_range": (
                    1,
                    len(paragraphs),
                ),
            }

        else:
            chunk_text = (
                "Cloud computing provides computing "
                "resources on demand."
            )
            source_location = {
                "page": 1,
            }

        chunk = Chunk(
            id=f"{document.document_id}:0",
            document_id=document.document_id,
            subject_id=document.subject_id,
            chunk_index=0,
            source=document.filename,
            text=chunk_text,
            **source_location,
        )

        embeddings = np.array(
            [[1.0, 0.0, 0.0]],
            dtype=np.float32,
        )

        self.vector_store.add(
            [chunk],
            embeddings,
        )
        self.document_repo.mark_indexed(
            document.document_id
        )

        return 1



class LifecycleQueryEmbedder:
    def embed_query(self, question):
        assert question == "Cloud computing là gì?"

        return np.array(
            [1.0, 0.0, 0.0],
            dtype=np.float32,
        )


class LifecycleAnswerGenerator:
    def generate(self, question, contexts):
        assert question == "Cloud computing là gì?"
        assert (
                "lesson.pdf" in contexts[0]
                or "lesson.txt" in contexts[0]
                or "lesson.docx" in contexts[0]
        )
        assert "computing resources on demand" in contexts[0]

        return (
            "Cloud computing cung cấp "
            "tài nguyên điện toán theo nhu cầu."
        )


@pytest.fixture
def client(tmp_path, monkeypatch):
    chroma_dir = tmp_path / "chroma"

    monkeypatch.setattr(
        config,
        "CHROMA_DIR",
        chroma_dir,
    )

    vector_store = VectorStore()

    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    Base.metadata.create_all(engine)

    TestingSessionLocal = sessionmaker(
        autocommit=False,
        autoflush=False,
        bind=engine,
    )
    db_session = TestingSessionLocal()

    document_repo = DocumentRepository(db_session)

    source_storage = LocalSourceStorage(
        root=tmp_path
    )

    pdf_service = LifecyclePDFService(
        document_repo,
        vector_store,
        source_storage,
    )

    chat_service = ChatService(
        embedder=LifecycleQueryEmbedder(),
        retriever=Retriever(vector_store),
        generator=LifecycleAnswerGenerator(),
    )

    deletion_service = DocumentDeletionService(
        document_repo=document_repo,
        vector_store=vector_store,
    )

    app = FastAPI()

    app.state.vector_store = vector_store
    app.state.document_repo = document_repo

    app.include_router(subjects_router)
    app.include_router(upload_router)
    app.include_router(chat_router)
    app.include_router(documents_router)

    app.dependency_overrides[
        get_source_storage
    ] = lambda: source_storage

    app.dependency_overrides[get_pdf_service] = (
        lambda: pdf_service
    )

    app.dependency_overrides[get_chat_service] = (
        lambda: chat_service
    )

    app.dependency_overrides[
        get_document_deletion_service
    ] = lambda: deletion_service

    app.dependency_overrides[
        get_document_repository
    ] = lambda: document_repo

    app.dependency_overrides[get_upload_dir] = (
        lambda: tmp_path
    )

    app.dependency_overrides[get_subject_repository] = (
        lambda: SubjectRepository(db_session)
    )

    app.dependency_overrides[
        get_upload_subject_repository
    ] = lambda: SubjectRepository(
        db_session
    )

    app.dependency_overrides[
        get_chat_subject_repository
    ] = lambda: SubjectRepository(
        db_session
    )

    try:
        with TestClient(app) as client:
            yield client
    finally:
        db_session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()



def test_create_subject_is_visible_in_subject_list(client):
    created = client.post(
        "/subjects/",
        json={"name": "Cloud Computing"},
    )

    assert created.status_code == 201
    subject = created.json()
    assert subject["subject_id"]

    listed = client.get("/subjects/")

    assert listed.status_code == 200
    assert listed.json() == [subject]


def test_vector_store_uses_temporary_chroma(
    client,
    tmp_path,
):
    assert config.CHROMA_DIR == tmp_path / "chroma"
    assert client.app.state.vector_store.collection.count() == 0


def test_document_lifecycle_upload_query_and_delete(client):
    pdf_data = (
        b"%PDF-1.4\ncontrolled test content"
    )

    subject_response = client.post(
        "/subjects/",
        json={"name": "Cloud Computing"},
    )
    subject_id = subject_response.json()["subject_id"]

    upload_response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": subject_id,
            "chapter": "1",
        },
    )

    assert upload_response.status_code == 201

    payload = upload_response.json()
    assert payload["filename"] == "lesson.pdf"
    assert payload["chunks"] == 1

    document = (
        client.app.state.document_repo.get_by_id(
            payload["document_id"]
        )
    )

    assert document.subject_id == subject_id
    assert document.status == "INDEXED"
    assert document.indexed_at is not None
    assert document.source_uri is not None
    assert Path(document.source_uri).exists()
    assert Path(document.source_uri).name == (
        f"{payload['document_id']}_lesson.pdf"
    )
    assert document.content_sha256 == (
        sha256(pdf_data).hexdigest()
    )
    assert document.file_size_bytes == len(pdf_data)

    indexed = (
        client.app.state.vector_store.collection.get(
            where={
                "document_id": payload["document_id"],
            }
        )
    )

    assert len(indexed["ids"]) == 1
    assert (
        indexed["metadatas"][0]["subject_id"]
        == subject_id
    )

    chat_response = client.post(
        "/chat/",
        json={
            "question": "Cloud computing là gì?",
            "subject_id": subject_id,
        },
    )

    assert chat_response.status_code == 200

    chat_payload = chat_response.json()

    assert chat_payload["answer"] == (
        "Cloud computing cung cấp "
        "tài nguyên điện toán theo nhu cầu."
    )
    assert len(chat_payload["sources"]) == 1
    assert (
        chat_payload["sources"][0]["document_id"]
        == payload["document_id"]
    )
    assert (
        chat_payload["sources"][0]["source"]
        == "lesson.pdf"
    )
    assert chat_payload["sources"][0]["page"] == 1

    delete_response = client.delete(
        f"/documents/{payload['document_id']}"
    )

    assert delete_response.status_code == 200
    assert delete_response.json() == {
        "document_id": payload["document_id"],
        "status": "DELETED",
        "deleted_chunks": 1,
    }

    deleted_document = (
        client.app.state.document_repo.get_by_id(
            payload["document_id"]
        )
    )

    assert deleted_document.status == "DELETED"
    assert deleted_document.deleted_at is not None

    remaining_chunks = (
        client.app.state.vector_store.collection.get(
            where={
                "document_id": payload["document_id"],
            }
        )
    )

    assert remaining_chunks["ids"] == []

    document_response = client.get(
        f"/documents/{payload['document_id']}"
    )

    assert document_response.status_code == 404

    chat_after_delete = client.post(
        "/chat/",
        json={
            "question": "Cloud computing là gì?",
            "subject_id": subject_id,
        },
    )

    assert chat_after_delete.status_code == 200
    assert chat_after_delete.json() == {
        "question": "Cloud computing là gì?",
        "answer": (
            "Tôi không tìm thấy thông tin "
            "trong tài liệu."
        ),
        "sources": [],
    }


def test_txt_lifecycle_upload_query_and_delete(
    client,
):
    text_data = (
        "Cloud computing provides computing "
        "resources on demand.\n"
        "Lineage is retained for every line."
    ).encode("utf-8")

    subject_response = client.post(
        "/subjects/",
        json={"name": "TXT Lifecycle"},
    )
    subject_id = subject_response.json()[
        "subject_id"
    ]

    upload_response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.txt",
                text_data,
                "text/plain",
            )
        },
        data={"subject_id": subject_id},
    )

    assert upload_response.status_code == 201

    payload = upload_response.json()
    document_id = payload["document_id"]

    document = (
        client.app.state.document_repo.get_by_id(
            document_id
        )
    )

    assert document.filename == "lesson.txt"
    assert document.file_type == "txt"
    assert document.status == "INDEXED"
    assert document.content_sha256 == (
        sha256(text_data).hexdigest()
    )
    assert document.file_size_bytes == len(
        text_data
    )

    indexed = (
        client.app.state.vector_store.collection.get(
            where={"document_id": document_id}
        )
    )

    assert len(indexed["ids"]) == 1

    metadata = indexed["metadatas"][0]

    assert metadata["source"] == "lesson.txt"
    assert metadata["line_start"] == 1
    assert metadata["line_end"] == 2
    assert "page" not in metadata
    assert "paragraph_start" not in metadata
    assert "paragraph_end" not in metadata

    chat_response = client.post(
        "/chat/",
        json={
            "question": "Cloud computing là gì?",
            "subject_id": subject_id,
        },
    )

    assert chat_response.status_code == 200

    source = chat_response.json()["sources"][0]

    assert source["document_id"] == document_id
    assert source["source"] == "lesson.txt"
    assert source["page"] is None
    assert source["line_range"] == [1, 2]
    assert source["paragraph_range"] is None

    delete_response = client.delete(
        f"/documents/{document_id}"
    )

    assert delete_response.status_code == 200
    assert delete_response.json()["status"] == (
        "DELETED"
    )
    assert (
        client.app.state.vector_store.collection.get(
            where={"document_id": document_id}
        )["ids"]
        == []
    )


def test_docx_lifecycle_upload_query_and_delete(
    client,
):
    buffer = BytesIO()

    word_document = WordDocument()
    word_document.add_paragraph(
        "Cloud computing provides computing "
        "resources on demand."
    )
    word_document.add_paragraph(
        "Paragraph lineage is retained."
    )
    word_document.save(buffer)

    docx_data = buffer.getvalue()

    subject_response = client.post(
        "/subjects/",
        json={"name": "DOCX Lifecycle"},
    )
    subject_id = subject_response.json()[
        "subject_id"
    ]

    upload_response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.docx",
                docx_data,
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml."
                    "document"
                ),
            )
        },
        data={"subject_id": subject_id},
    )

    assert upload_response.status_code == 201

    payload = upload_response.json()
    document_id = payload["document_id"]

    document = (
        client.app.state.document_repo.get_by_id(
            document_id
        )
    )

    assert document.filename == "lesson.docx"
    assert document.file_type == "docx"
    assert document.status == "INDEXED"
    assert document.content_sha256 == (
        sha256(docx_data).hexdigest()
    )
    assert document.file_size_bytes == len(
        docx_data
    )

    indexed = (
        client.app.state.vector_store.collection.get(
            where={"document_id": document_id}
        )
    )

    assert len(indexed["ids"]) == 1

    metadata = indexed["metadatas"][0]

    assert metadata["source"] == "lesson.docx"
    assert metadata["paragraph_start"] == 1
    assert metadata["paragraph_end"] == 2
    assert "page" not in metadata
    assert "line_start" not in metadata
    assert "line_end" not in metadata

    chat_response = client.post(
        "/chat/",
        json={
            "question": "Cloud computing là gì?",
            "subject_id": subject_id,
        },
    )

    assert chat_response.status_code == 200

    source = chat_response.json()["sources"][0]

    assert source["document_id"] == document_id
    assert source["source"] == "lesson.docx"
    assert source["page"] is None
    assert source["line_range"] is None
    assert source["paragraph_range"] == [
        1,
        2,
    ]

    delete_response = client.delete(
        f"/documents/{document_id}"
    )

    assert delete_response.status_code == 200
    assert delete_response.json()["status"] == (
        "DELETED"
    )
    assert (
        client.app.state.vector_store.collection.get(
            where={"document_id": document_id}
        )["ids"]
        == []
    )

def test_repeated_upload_does_not_create_duplicate_effects(
    client,
    tmp_path,
):
    pdf_data = b"%PDF-1.4\nsame document content"

    subject_response = client.post(
        "/subjects/",
        json={"name": "Cloud Computing"},
    )
    subject_id = subject_response.json()["subject_id"]

    first_response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": subject_id,
            "chapter": "1",
        },
    )

    assert first_response.status_code == 201
    first_payload = first_response.json()

    repeated_response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-copy.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": subject_id,
            "chapter": "1",
        },
    )

    assert repeated_response.status_code == 409
    assert repeated_response.json() == {
        "detail": {
            "code": "DOCUMENT_ALREADY_EXISTS",
            "document_id": first_payload["document_id"],
            "status": "INDEXED",
        }
    }

    active_documents = (
        client.app.state.document_repo.list_active()
    )

    assert len(active_documents) == 1
    assert (
        active_documents[0].document_id
        == first_payload["document_id"]
    )

    stored_pdfs = list(tmp_path.glob("*.pdf"))

    assert len(stored_pdfs) == 1
    assert stored_pdfs[0].name == (
        f"{first_payload['document_id']}_lesson.pdf"
    )

    assert (
        client.app.state.vector_store.collection.count()
        == 1
    )


def test_failed_document_retry_reuses_source_and_identity(
    client,
    tmp_path,
):
    pdf_data = b"%PDF-1.4\nretryable lifecycle content"

    subject_response = client.post(
        "/subjects/",
        json={"name": "Cloud Computing"},
    )
    subject_id = subject_response.json()["subject_id"]

    document_id = str(uuid4())
    source_path = (
        tmp_path / f"{document_id}_lesson.pdf"
    )
    source_path.write_bytes(pdf_data)

    repository = client.app.state.document_repo

    repository.create(
        document_id=document_id,
        subject_id=subject_id,
        filename="lesson.pdf",
        file_type="pdf",
        chapter=1,
        source_uri=source_path.as_posix(),
        content_sha256=sha256(
            pdf_data
        ).hexdigest(),
        file_size_bytes=len(pdf_data),
    )

    repository.mark_failed(
        document_id=document_id,
        failed_stage="EXTRACT",
        error_message="Controlled initial failure",
    )

    response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-retry.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": subject_id,
            "chapter": "1",
        },
    )

    assert response.status_code == 200
    assert response.json() == {
        "document_id": document_id,
        "filename": "lesson.pdf",
        "chunks": 1,
    }

    documents = repository.list_active()

    assert len(documents) == 1
    assert documents[0].document_id == document_id
    assert documents[0].status == "INDEXED"
    assert documents[0].failed_stage is None
    assert documents[0].error_message is None
    assert documents[0].source_uri == (
        source_path.as_posix()
    )

    stored_pdfs = list(tmp_path.glob("*.pdf"))

    assert stored_pdfs == [source_path]

    indexed = (
        client.app.state.vector_store.collection.get(
            where={"document_id": document_id}
        )
    )

    assert len(indexed["ids"]) == 1


def test_repeated_txt_upload_does_not_duplicate_effects(
    client,
    tmp_path,
):
    text_data = (
        b"first line\n"
        b"second line"
    )

    subject_response = client.post(
        "/subjects/",
        json={"name": "TXT Duplicate"},
    )
    subject_id = subject_response.json()[
        "subject_id"
    ]

    first_response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.txt",
                text_data,
                "text/plain",
            )
        },
        data={"subject_id": subject_id},
    )

    assert first_response.status_code == 201
    first_payload = first_response.json()

    repeated_response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-copy.txt",
                text_data,
                "text/plain",
            )
        },
        data={"subject_id": subject_id},
    )

    assert repeated_response.status_code == 409
    assert repeated_response.json() == {
        "detail": {
            "code": "DOCUMENT_ALREADY_EXISTS",
            "document_id": (
                first_payload["document_id"]
            ),
            "status": "INDEXED",
        }
    }

    documents = (
        client.app.state.document_repo.list_active()
    )

    assert len(documents) == 1
    assert documents[0].file_type == "txt"

    stored_sources = list(
        tmp_path.glob("*.txt")
    )

    assert len(stored_sources) == 1
    assert stored_sources[0].name == (
        f"{first_payload['document_id']}"
        "_lesson.txt"
    )

    assert (
        client.app.state
        .vector_store.collection.count()
        == 1
    )


def test_failed_txt_retry_reuses_source_and_identity(
    client,
    tmp_path,
):
    text_data = (
        b"first line\n"
        b"second line"
    )

    subject_response = client.post(
        "/subjects/",
        json={"name": "TXT Retry"},
    )
    subject_id = subject_response.json()[
        "subject_id"
    ]

    document_id = str(uuid4())
    source_path = (
        tmp_path
        / f"{document_id}_lesson.txt"
    )
    source_path.write_bytes(text_data)

    repository = client.app.state.document_repo

    repository.create(
        document_id=document_id,
        subject_id=subject_id,
        filename="lesson.txt",
        file_type="txt",
        chapter=None,
        source_uri=source_path.as_posix(),
        content_sha256=sha256(
            text_data
        ).hexdigest(),
        file_size_bytes=len(text_data),
    )
    repository.mark_failed(
        document_id=document_id,
        failed_stage="EXTRACT",
        error_message="Controlled TXT failure",
    )

    response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-retry.txt",
                text_data,
                "text/plain",
            )
        },
        data={"subject_id": subject_id},
    )

    assert response.status_code == 200
    assert response.json() == {
        "document_id": document_id,
        "filename": "lesson.txt",
        "chunks": 1,
    }

    documents = repository.list_active()

    assert len(documents) == 1
    assert documents[0].document_id == (
        document_id
    )
    assert documents[0].file_type == "txt"
    assert documents[0].status == "INDEXED"
    assert documents[0].failed_stage is None
    assert documents[0].error_message is None
    assert documents[0].source_uri == (
        source_path.as_posix()
    )

    assert list(tmp_path.glob("*.txt")) == [
        source_path
    ]

    indexed = (
        client.app.state.vector_store.collection.get(
            where={"document_id": document_id}
        )
    )

    assert len(indexed["ids"]) == 1

    metadata = indexed["metadatas"][0]

    assert metadata["line_start"] == 1
    assert metadata["line_end"] == 2
    assert "page" not in metadata


def test_repeated_docx_upload_does_not_duplicate_effects(
    client,
    tmp_path,
):
    buffer = BytesIO()

    word_document = WordDocument()
    word_document.add_paragraph(
        "Duplicate DOCX content"
    )
    word_document.save(buffer)

    docx_data = buffer.getvalue()
    content_type = (
        "application/vnd.openxmlformats-"
        "officedocument.wordprocessingml."
        "document"
    )

    subject_response = client.post(
        "/subjects/",
        json={"name": "DOCX Duplicate"},
    )
    subject_id = subject_response.json()[
        "subject_id"
    ]

    first_response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.docx",
                docx_data,
                content_type,
            )
        },
        data={"subject_id": subject_id},
    )

    assert first_response.status_code == 201
    first_payload = first_response.json()

    repeated_response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-copy.docx",
                docx_data,
                content_type,
            )
        },
        data={"subject_id": subject_id},
    )

    assert repeated_response.status_code == 409
    assert repeated_response.json() == {
        "detail": {
            "code": "DOCUMENT_ALREADY_EXISTS",
            "document_id": (
                first_payload["document_id"]
            ),
            "status": "INDEXED",
        }
    }

    documents = (
        client.app.state.document_repo.list_active()
    )

    assert len(documents) == 1
    assert documents[0].file_type == "docx"

    stored_sources = list(
        tmp_path.glob("*.docx")
    )

    assert len(stored_sources) == 1
    assert stored_sources[0].name == (
        f"{first_payload['document_id']}"
        "_lesson.docx"
    )

    assert (
        client.app.state
        .vector_store.collection.count()
        == 1
    )


def test_failed_docx_retry_reuses_source_and_identity(
    client,
    tmp_path,
):
    buffer = BytesIO()

    word_document = WordDocument()
    word_document.add_paragraph(
        "First paragraph"
    )
    word_document.add_paragraph(
        "Second paragraph"
    )
    word_document.save(buffer)

    docx_data = buffer.getvalue()
    content_type = (
        "application/vnd.openxmlformats-"
        "officedocument.wordprocessingml."
        "document"
    )

    subject_response = client.post(
        "/subjects/",
        json={"name": "DOCX Retry"},
    )
    subject_id = subject_response.json()[
        "subject_id"
    ]

    document_id = str(uuid4())
    source_path = (
        tmp_path
        / f"{document_id}_lesson.docx"
    )
    source_path.write_bytes(docx_data)

    repository = client.app.state.document_repo

    repository.create(
        document_id=document_id,
        subject_id=subject_id,
        filename="lesson.docx",
        file_type="docx",
        chapter=None,
        source_uri=source_path.as_posix(),
        content_sha256=sha256(
            docx_data
        ).hexdigest(),
        file_size_bytes=len(docx_data),
    )
    repository.mark_failed(
        document_id=document_id,
        failed_stage="EXTRACT",
        error_message="Controlled DOCX failure",
    )

    response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-retry.docx",
                docx_data,
                content_type,
            )
        },
        data={"subject_id": subject_id},
    )

    assert response.status_code == 200
    assert response.json() == {
        "document_id": document_id,
        "filename": "lesson.docx",
        "chunks": 1,
    }

    documents = repository.list_active()

    assert len(documents) == 1
    assert documents[0].document_id == (
        document_id
    )
    assert documents[0].file_type == "docx"
    assert documents[0].status == "INDEXED"
    assert documents[0].failed_stage is None
    assert documents[0].error_message is None
    assert documents[0].source_uri == (
        source_path.as_posix()
    )

    assert list(tmp_path.glob("*.docx")) == [
        source_path
    ]

    indexed = (
        client.app.state.vector_store.collection.get(
            where={"document_id": document_id}
        )
    )

    assert len(indexed["ids"]) == 1

    metadata = indexed["metadatas"][0]

    assert metadata["paragraph_start"] == 1
    assert metadata["paragraph_end"] == 2
    assert "page" not in metadata
    assert "line_start" not in metadata
