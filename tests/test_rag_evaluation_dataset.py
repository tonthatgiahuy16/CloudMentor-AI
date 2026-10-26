import json
from pathlib import Path

from copy import deepcopy

import pytest

from evaluation.dataset import validate_dataset

from docx import Document

from pypdf import PdfWriter

@pytest.fixture
def valid_dataset(tmp_path):
    source_path = tmp_path / "source.txt"
    source_path.write_text(
        "First line.\nEvidence line.\n",
        encoding="utf-8",
    )

    return {
        "schema_version": "1.0",
        "sources": [
            {
                "source_id": "source-1",
                "subject_key": "subject-1",
                "path": "source.txt",
                "format": "txt",
            }
        ],
        "cases": [
            {
                "case_id": "case-1",
                "subject_key": "subject-1",
                "question": "What is on the second line?",
                "answerable": True,
                "expected_evidence": [
                    {
                        "source_id": "source-1",
                        "line_range": [2, 2],
                    }
                ],
                "expected_facts": ["Evidence line."],
            }
        ],
    }


def test_accepts_valid_dataset(valid_dataset, tmp_path):
    validate_dataset(valid_dataset, tmp_path)


def test_rejects_unknown_evidence_source(valid_dataset, tmp_path):
    dataset = deepcopy(valid_dataset)
    dataset["cases"][0]["expected_evidence"][0]["source_id"] = "missing"

    with pytest.raises(ValueError, match="unknown source_id"):
        validate_dataset(dataset, tmp_path)


def test_rejects_evidence_outside_source(valid_dataset, tmp_path):
    dataset = deepcopy(valid_dataset)
    dataset["cases"][0]["expected_evidence"][0]["line_range"] = [3, 3]

    with pytest.raises(ValueError, match="out of bounds"):
        validate_dataset(dataset, tmp_path)


@pytest.fixture
def docx_dataset(valid_dataset, tmp_path):
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Evidence paragraph.")
    document.save(tmp_path / "source.docx")

    dataset = deepcopy(valid_dataset)
    dataset["sources"][0]["path"] = "source.docx"
    dataset["sources"][0]["format"] = "docx"
    dataset["cases"][0]["expected_evidence"][0] = {
        "source_id": "source-1",
        "paragraph_range": [2, 2],
    }

    return dataset


def test_accepts_valid_docx_evidence(docx_dataset, tmp_path):
    validate_dataset(docx_dataset, tmp_path)


def test_rejects_paragraph_outside_docx(docx_dataset, tmp_path):
    dataset = deepcopy(docx_dataset)
    dataset["cases"][0]["expected_evidence"][0]["paragraph_range"] = [3, 3]

    with pytest.raises(ValueError, match="paragraph_range out of bounds"):
        validate_dataset(dataset, tmp_path)


@pytest.fixture
def pdf_dataset(valid_dataset, tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)

    with (tmp_path / "source.pdf").open("wb") as stream:
        writer.write(stream)

    dataset = deepcopy(valid_dataset)
    dataset["sources"][0]["path"] = "source.pdf"
    dataset["sources"][0]["format"] = "pdf"
    dataset["cases"][0]["expected_evidence"][0] = {
        "source_id": "source-1",
        "page": 1,
    }

    return dataset


def test_accepts_valid_pdf_evidence(pdf_dataset, tmp_path):
    validate_dataset(pdf_dataset, tmp_path)


def test_rejects_page_outside_pdf(pdf_dataset, tmp_path):
    dataset = deepcopy(pdf_dataset)
    dataset["cases"][0]["expected_evidence"][0]["page"] = 2

    with pytest.raises(ValueError, match="page out of bounds"):
        validate_dataset(dataset, tmp_path)


def test_project_evaluation_dataset_is_valid():
    dataset_path = (
        Path(__file__).resolve().parents[1]
        / "evaluation"
        / "cases.json"
    )

    dataset = json.loads(
        dataset_path.read_text(encoding="utf-8-sig")
    )

    validate_dataset(dataset, dataset_path.parent)
