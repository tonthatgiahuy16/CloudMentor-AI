
import pytest
from io import BytesIO
from app.rag import loader as loader_module
from docx import Document as WordDocument
from app.rag.loader import (
    DOCXLoader,
    PDFLoader,
    TextLoader,
)


class FakePage:
    def __init__(self, text):
        self.text = text

    def extract_text(self):
        return self.text


class FakeReader:
    def __init__(self, _path):
        self.pages = [FakePage("first page"), FakePage(None)]


def test_pdf_loader_extracts_pages_from_bytes(
    monkeypatch,
):
    source_data = b"%PDF-1.4"

    monkeypatch.setattr(
        loader_module,
        "PdfReader",
        FakeReader,
    )

    documents = PDFLoader().load(
        source_data,
        filename="lesson.pdf",
        document_id="doc-1",
        subject_id="subject-1",
    )

    assert [
        document.text
        for document in documents
    ] == ["first page", ""]

    assert documents[0].source == "lesson.pdf"
    assert documents[0].page == 1
    assert documents[0].document_id == "doc-1"
    assert documents[0].subject_id == "subject-1"


def test_loader_rejects_non_pdf():
    with pytest.raises(
        ValueError,
        match="PDF",
    ):
        PDFLoader().load(
            b"not a PDF source",
            filename="lesson.txt",
            document_id="doc-1",
            subject_id="subject-1",
        )


def test_text_loader_preserves_exact_line_locations():
    source_data = (
        b"first line\r\n\r\nthird line\r\n"
    )

    documents = TextLoader().load(
        source_data,
        filename="lesson.txt",
        document_id="doc-1",
        subject_id="subject-1",
    )

    assert len(documents) == 1

    document = documents[0]

    assert document.document_id == "doc-1"
    assert document.subject_id == "subject-1"
    assert document.source == "lesson.txt"
    assert document.text == (
        "first line\n\nthird line"
    )
    assert document.page is None
    assert document.line_range == (1, 3)
    assert document.paragraph_range is None


def test_text_loader_rejects_non_txt():
    with pytest.raises(
        ValueError,
        match="TXT",
    ):
        TextLoader().load(
            b"not a text source",
            filename="lesson.pdf",
            document_id="doc-1",
            subject_id="subject-1",
        )


def test_text_loader_rejects_empty_content():
    with pytest.raises(
        ValueError,
        match="empty",
    ):
        TextLoader().load(
            b" \r\n\t\r\n",
            filename="empty.txt",
            document_id="doc-1",
            subject_id="subject-1",
        )


def test_text_loader_rejects_invalid_utf8():
    with pytest.raises(
        ValueError,
        match="valid UTF-8",
    ):
        TextLoader().load(
            b"\xff\xfe\xfa",
            filename="invalid.txt",
            document_id="doc-1",
            subject_id="subject-1",
        )


def test_docx_loader_preserves_exact_paragraph_locations():
    buffer = BytesIO()

    word_document = WordDocument()
    word_document.add_paragraph(
        "First paragraph"
    )
    word_document.add_paragraph("")
    word_document.add_paragraph(
        "Third paragraph"
    )
    word_document.save(buffer)

    documents = DOCXLoader().load(
        buffer.getvalue(),
        filename="lesson.docx",
        document_id="doc-1",
        subject_id="subject-1",
    )

    assert len(documents) == 1

    document = documents[0]

    assert document.document_id == "doc-1"
    assert document.subject_id == "subject-1"
    assert document.source == "lesson.docx"
    assert document.text == (
        "First paragraph\n"
        "\n"
        "Third paragraph"
    )
    assert document.page is None
    assert document.line_range is None
    assert document.paragraph_range == (
        1,
        3,
    )


def test_docx_loader_rejects_nonempty_tables():
    buffer = BytesIO()
    word_document = WordDocument()

    word_document.add_paragraph(
        "Body paragraph"
    )

    table = word_document.add_table(
        rows=1,
        cols=2,
    )
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"

    word_document.save(buffer)

    with pytest.raises(
        ValueError,
        match="DOCX tables are not supported",
    ):
        DOCXLoader().load(
            buffer.getvalue(),
            filename="with-table.docx",
            document_id="doc-1",
            subject_id="subject-1",
        )


def test_docx_loader_rejects_empty_document():
    buffer = BytesIO()
    WordDocument().save(buffer)

    with pytest.raises(
        ValueError,
        match="DOCX file is empty",
    ):
        DOCXLoader().load(
            buffer.getvalue(),
            filename="empty.docx",
            document_id="doc-1",
            subject_id="subject-1",
        )


def test_docx_loader_rejects_invalid_container():
    with pytest.raises(
        ValueError,
        match="DOCX file is not valid",
    ):
        DOCXLoader().load(
            b"not a valid DOCX package",
            filename="invalid.docx",
            document_id="doc-1",
            subject_id="subject-1",
        )


def test_docx_loader_normalizes_internal_line_breaks():
    buffer = BytesIO()
    word_document = WordDocument()

    first_paragraph = (
        word_document.add_paragraph(
            "First line"
        )
    )
    first_paragraph.add_run().add_break()
    first_paragraph.add_run(
        "same paragraph"
    )

    word_document.add_paragraph(
        "Second paragraph"
    )
    word_document.save(buffer)

    documents = DOCXLoader().load(
        buffer.getvalue(),
        filename="line-breaks.docx",
        document_id="doc-1",
        subject_id="subject-1",
    )

    assert documents[0].text == (
        "First line same paragraph\n"
        "Second paragraph"
    )
    assert (
        documents[0].paragraph_range
        == (1, 2)
    )
