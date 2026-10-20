from hashlib import sha256
from types import SimpleNamespace
from unittest.mock import Mock
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import app.api.upload as upload_api
from app.api.upload import (
    get_pdf_service,
    get_upload_dir,
    get_upload_subject_repository,
    get_source_storage,
    router,
)
from app.repos.document_repository import (
    ActiveDocumentConflictError,
)
from app.storage.source_storage import (
    LocalSourceStorage,
)

from io import BytesIO

from docx import Document as WordDocument



class FakeSubjectRepository:
    def __init__(self, subject):
        self.subject = subject
        self.calls = []

    def get_by_id(self, subject_id):
        self.calls.append(subject_id)
        return self.subject


class FakePDFService:
    def __init__(self):
        self.calls = []

    def find_active_duplicate(
        self,
        *,
        subject_id,
        content_sha256,
    ):
        return None

    def upload(
        self,
        *,
        source_data,
        subject_id,
        chapter,
        document_id,
        original_filename,
        content_sha256=None,
        file_size_bytes=None,
    ):
        self.calls.append(
            (
                source_data,
                subject_id,
                chapter,
                document_id,
                original_filename,
                content_sha256,
                file_size_bytes,
            )
        )
        return 7


class DuplicatePDFService(FakePDFService):
    def find_active_duplicate(
        self,
        *,
        subject_id,
        content_sha256,
    ):
        return SimpleNamespace(
            document_id="existing-document-id",
            filename="existing.pdf",
            status="INDEXED",
        )


class ConcurrentDuplicatePDFService(FakePDFService):
    def __init__(self):
        super().__init__()
        self.existing_document = SimpleNamespace(
            document_id="winning-document-id",
            filename="original.pdf",
            status="PROCESSING",
        )
        self.upload_calls = []

    def upload(self, **kwargs):
        self.upload_calls.append(kwargs)
        raise ActiveDocumentConflictError(
            self.existing_document
        )


class FailedDocumentPDFService(FakePDFService):
    def __init__(self):
        super().__init__()
        self.existing_document = SimpleNamespace(
            document_id="failed-document-id",
            filename="original.pdf",
            status="FAILED",
            source_uri="storage/uploads/failed-document.pdf",
        )
        self.retry_calls = []

    def find_active_duplicate(
        self,
        *,
        subject_id,
        content_sha256,
    ):
        return self.existing_document

    def retry_failed_document(
        self,
        document,
    ):
        self.retry_calls.append(document)
        return 5


class FailingPDFService:
    def upload(self, *args, **kwargs):
        raise RuntimeError(
            "Controlled processing failure"
        )

    def find_active_duplicate(
        self,
        *,
        subject_id,
        content_sha256,
    ):
        return None


class FailingRetryPDFService(
    FailedDocumentPDFService
):
    def retry_failed_document(
        self,
        document,
    ):
        self.retry_calls.append(document)
        raise RuntimeError(
            "Controlled retry failure"
        )

DEFAULT_SUBJECT = object()

def make_client(
    tmp_path,
    service=None,
    subject=DEFAULT_SUBJECT,
    source_storage=None,
):
    if service is None:
        service = FakePDFService()

    if source_storage is None:
        source_storage = LocalSourceStorage(
            root=tmp_path,
        )

    if subject is DEFAULT_SUBJECT:
        subject = SimpleNamespace(
            subject_id="subject-active",
            status="ACTIVE",
        )

    subject_repo = FakeSubjectRepository(
        subject
    )

    app = FastAPI()
    app.include_router(router)

    app.dependency_overrides[get_pdf_service] = (
        lambda: service
    )
    app.dependency_overrides[get_upload_dir] = (
        lambda: tmp_path
    )
    app.dependency_overrides[get_source_storage] = (
        lambda: source_storage
    )
    app.dependency_overrides[
        get_upload_subject_repository
    ] = lambda: subject_repo

    return TestClient(app), service


def test_upload_accepts_pdf_and_passes_source_bytes(
    tmp_path,
):
    client, service = make_client(tmp_path)

    pdf_data = b"%PDF-1.4\ncontent"

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
            "chapter": "3",
        },
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["filename"] == "lesson.pdf"
    assert payload["chunks"] == 7
    (
        source_data,
        subject_id,
        chapter,
        document_id,
        original_filename,
        content_sha256,
        file_size_bytes,
    ) = service.calls[0]
    assert subject_id == "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
    assert chapter == 3
    assert original_filename == "lesson.pdf"
    assert source_data == pdf_data
    assert content_sha256 == sha256(pdf_data).hexdigest()
    assert file_size_bytes == len(pdf_data)


def test_upload_rejects_fake_pdf_content(tmp_path):
    client, _ = make_client(tmp_path)

    response = client.post(
        "/upload/",
        files={"file": ("lesson.pdf", b"not a pdf", "application/pdf")},
        data={"subject_id": "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"},
    )

    assert response.status_code == 415
    assert not list(tmp_path.iterdir())


def test_upload_rejects_wrong_extension(tmp_path):
    client, _ = make_client(tmp_path)

    response = client.post(
        "/upload/",
        files={"file": ("lesson.txt", b"%PDF-1.4", "application/pdf")},
        data={"subject_id": "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"},
    )

    assert response.status_code == 415


def test_upload_returns_500_when_processing_fails(
    tmp_path,
):
    service = FailingPDFService()
    client, _ = make_client(
        tmp_path,
        service=service,
    )

    pdf_data = b"%PDF-1.4\nrecoverable content"

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == 500

    assert response.json() == {
    "detail": "Document processing failed"
    }


def test_upload_rejects_active_duplicate_before_processing(
    tmp_path,
):
    service = DuplicatePDFService()
    client, _ = make_client(
        tmp_path,
        service=service,
    )

    pdf_data = b"%PDF-1.4\nduplicate content"

    response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-copy.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == 409
    assert response.json() == {
        "detail": {
            "code": "DOCUMENT_ALREADY_EXISTS",
            "document_id": "existing-document-id",
            "status": "INDEXED",
        }
    }
    assert service.calls == []
    assert list(tmp_path.iterdir()) == []


def test_upload_returns_409_when_database_reservation_conflicts(
    tmp_path,
):
    service = ConcurrentDuplicatePDFService()
    client, _ = make_client(
        tmp_path,
        service=service,
    )

    pdf_data = b"%PDF-1.4\nconcurrent duplicate"

    response = client.post(
        "/upload/",
        files={
            "file": (
                "same-content.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == 409
    assert response.json() == {
        "detail": {
            "code": "DOCUMENT_ALREADY_EXISTS",
            "document_id": "winning-document-id",
            "status": "PROCESSING",
        }
    }
    assert service.upload_calls == [{
        "source_data": pdf_data,
        "subject_id": (
            "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
        ),
        "chapter": None,
        "document_id": service.upload_calls[0]["document_id"],
        "original_filename": "same-content.pdf",
        "content_sha256": sha256(pdf_data).hexdigest(),
        "file_size_bytes": len(pdf_data),
    }]
    assert not list(tmp_path.iterdir())


def test_upload_retries_failed_document_without_new_source_file(
    tmp_path,
):
    service = FailedDocumentPDFService()
    client, _ = make_client(
        tmp_path,
        service=service,
    )

    pdf_data = b"%PDF-1.4\nretryable content"

    response = client.post(
        "/upload/",
        files={
            "file": (
                "renamed-retry.pdf",
                pdf_data,
                "application/pdf",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == 200
    assert response.json() == {
        "document_id": "failed-document-id",
        "filename": "original.pdf",
        "chunks": 5,
    }

    assert service.calls == []
    assert service.retry_calls == [
        service.existing_document
    ]
    assert list(tmp_path.iterdir()) == []


def test_upload_returns_500_when_failed_document_retry_fails(
    tmp_path,
):
    service = FailingRetryPDFService()
    client, _ = make_client(
        tmp_path,
        service=service,
    )

    response = client.post(
        "/upload/",
        files={
            "file": (
                "retry.pdf",
                b"%PDF-1.4\nretryable content",
                "application/pdf",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == 500
    assert response.json() == {
        "detail": "Document retry failed"
    }
    assert service.calls == []
    assert service.retry_calls == [
        service.existing_document
    ]
    assert list(tmp_path.iterdir()) == []


def test_upload_rejects_subject_in_trash_before_writing(
    tmp_path,
):
    subject = SimpleNamespace(
        subject_id="subject-deleted",
        status="DELETED",
    )
    client, service = make_client(
        tmp_path,
        subject=subject,
    )

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.pdf",
                b"%PDF-1.4\ncontent",
                "application/pdf",
            )
        },
        data={
            "subject_id": "subject-deleted"
        },
    )

    assert response.status_code == 409
    assert response.json() == {
        "detail": "Subject is not active"
    }
    assert service.calls == []
    assert not list(tmp_path.iterdir())


def test_upload_rejects_missing_subject_before_writing(
    tmp_path,
):
    client, service = make_client(
        tmp_path,
        subject=None,
    )

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.pdf",
                b"%PDF-1.4\ncontent",
                "application/pdf",
            )
        },
        data={
            "subject_id": "missing-subject"
        },
    )

    assert response.status_code == 404
    assert response.json() == {
        "detail": "Subject not found"
    }
    assert service.calls == []
    assert not list(tmp_path.iterdir())


def test_upload_accepts_txt_and_passes_source_bytes(
    tmp_path,
):
    client, service = make_client(tmp_path)

    text_data = (
        "first line\r\n"
        "second line\r\n"
    ).encode("utf-8")

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.txt",
                text_data,
                "text/plain",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == 201
    assert response.json()["filename"] == (
        "lesson.txt"
    )
    assert response.json()["chunks"] == 7

    (
        source_data,
        subject_id,
        chapter,
        document_id,
        original_filename,
        content_sha256,
        file_size_bytes,
    ) = service.calls[0]

    assert subject_id == (
        "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
    )
    assert chapter is None
    assert original_filename == "lesson.txt"
    assert source_data == text_data
    assert content_sha256 == (
        sha256(text_data).hexdigest()
    )
    assert file_size_bytes == len(text_data)


@pytest.mark.parametrize(
    ("data", "expected_status", "expected_detail"),
    [
        (
            b" \r\n\t\r\n",
            400,
            "TXT file is empty",
        ),
        (
            b"\xff\xfe\xfa",
            415,
            "TXT file is not valid UTF-8",
        ),
    ],
)
def test_upload_rejects_invalid_txt_before_writing(
    tmp_path,
    data,
    expected_status,
    expected_detail,
):
    client, service = make_client(tmp_path)

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.txt",
                data,
                "text/plain",
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == expected_status
    assert response.json() == {
        "detail": expected_detail,
    }
    assert service.calls == []
    assert not list(tmp_path.iterdir())


def test_upload_accepts_docx_and_passes_source_bytes(
    tmp_path,
):
    buffer = BytesIO()

    word_document = WordDocument()
    word_document.add_paragraph(
        "First paragraph"
    )
    word_document.add_paragraph(
        "Second paragraph"
    )
    word_document.save(buffer)

    docx_data = buffer.getvalue()

    client, service = make_client(tmp_path)

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.docx",
                docx_data,
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml."
                    "document"
                ),
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == 201
    assert response.json()["filename"] == (
        "lesson.docx"
    )
    assert response.json()["chunks"] == 7

    (
        source_data,
        subject_id,
        chapter,
        document_id,
        original_filename,
        content_sha256,
        file_size_bytes,
    ) = service.calls[0]

    assert original_filename == "lesson.docx"
    assert source_data == docx_data
    assert content_sha256 == (
        sha256(docx_data).hexdigest()
    )
    assert file_size_bytes == len(docx_data)


@pytest.mark.parametrize(
    (
        "case",
        "expected_status",
        "expected_detail",
    ),
    [
        (
            "empty",
            400,
            "DOCX file is empty",
        ),
        (
            "table",
            415,
            "DOCX tables are not supported",
        ),
        (
            "corrupt",
            415,
            "DOCX file is not valid",
        ),
    ],
)
def test_upload_rejects_invalid_docx_before_writing(
    tmp_path,
    case,
    expected_status,
    expected_detail,
):
    if case == "corrupt":
        docx_data = b"not a DOCX package"
    else:
        buffer = BytesIO()
        word_document = WordDocument()

        if case == "table":
            word_document.add_paragraph(
                "Body paragraph"
            )
            table = word_document.add_table(
                rows=1,
                cols=1,
            )
            table.cell(0, 0).text = (
                "Table content"
            )

        word_document.save(buffer)
        docx_data = buffer.getvalue()

    client, service = make_client(tmp_path)

    response = client.post(
        "/upload/",
        files={
            "file": (
                "lesson.docx",
                docx_data,
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml."
                    "document"
                ),
            )
        },
        data={
            "subject_id": (
                "7d13e6d0-ef38-4bf2-a5cf-a593517087e4"
            ),
        },
    )

    assert response.status_code == (
        expected_status
    )
    assert response.json() == {
        "detail": expected_detail,
    }
    assert service.calls == []
    assert not list(tmp_path.iterdir())


def test_get_pdf_service_injects_source_storage(
    monkeypatch,
):
    source_storage = Mock()
    db = Mock()
    document_repo = Mock()
    pdf_service = Mock()

    session_local = Mock(return_value=db)
    document_repository_class = Mock(
        return_value=document_repo
    )
    pdf_service_class = Mock(
        return_value=pdf_service
    )

    monkeypatch.setattr(
        "app.api.upload.SessionLocal",
        session_local,
    )
    monkeypatch.setattr(
        "app.api.upload.DocumentRepository",
        document_repository_class,
    )
    monkeypatch.setattr(
        "app.services.pdf_service.PDFService",
        pdf_service_class,
    )

    dependency = get_pdf_service(
        source_storage=source_storage
    )

    assert next(dependency) is pdf_service

    pdf_service_class.assert_called_once_with(
        document_repo,
        source_storage=source_storage,
    )

    dependency.close()
    db.close.assert_called_once_with()


def test_get_source_storage_uses_factory(
    monkeypatch,
):
    source_storage = Mock()
    storage_factory = Mock(
        return_value=source_storage
    )

    monkeypatch.setattr(
        upload_api,
        "create_source_storage",
        storage_factory,
        raising=False,
    )

    result = get_source_storage()

    assert result is source_storage
    storage_factory.assert_called_once_with()
