from hashlib import sha256
from inspect import signature
from types import SimpleNamespace
from unittest.mock import Mock, call
import pytest

from app.services.pdf_service import (
    PDFService,
    RetainedSourceValidationError,
)
from app.pipeline.chunk_pipeline import ChunkPipeline
from app.pipeline.extract_pipeline import ExtractPipeline
from app.pipeline.transform_pipeline import TransformPipeline
from app.rag.chunker import TextChunker
from app.storage.source_storage import (
    LocalSourceStorage,
    SourceUnavailableError,
    SourceNotFoundError,
    StoredSource,
)

from docx import Document as WordDocument


def test_retry_failed_document_reuses_source_and_document_id(
    tmp_path,
):
    pdf_data = b"%PDF-1.4\nretained source"
    source_path = tmp_path / "retained.pdf"
    source_path.write_bytes(pdf_data)

    document = SimpleNamespace(
        document_id="failed-document-id",
        subject_id="subject-id",
        filename="lesson.pdf",
        file_type="pdf",
        status="FAILED",
        source_uri=source_path.as_posix(),
        content_sha256=sha256(pdf_data).hexdigest(),
        file_size_bytes=len(pdf_data),
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = LocalSourceStorage(
        root=tmp_path,
    )

    service.document_repo = Mock()
    service.extract = Mock()
    service.transform = Mock()
    service.chunk = Mock()
    service.embedding = Mock()
    service.index = Mock()

    extracted = [object()]
    transformed = [object()]
    chunks = [object(), object()]
    embeddings = object()

    service.extract.run.return_value = extracted
    service.transform.run.return_value = transformed
    service.chunk.run.return_value = chunks
    service.embedding.run.return_value = embeddings

    result = service.retry_failed_document(document)

    assert result == 2

    assert (
        service.document_repo
        .record_processing_heartbeat
        .call_args_list
    ) == [
        call(
            document_id="failed-document-id",
            stage="EXTRACT",
        ),
        call(
            document_id="failed-document-id",
            stage="TRANSFORM",
        ),
        call(
            document_id="failed-document-id",
            stage="CHUNK",
        ),
        call(
            document_id="failed-document-id",
            stage="EMBEDDING",
        ),
        call(
            document_id="failed-document-id",
            stage="INDEX",
        ),
    ]

    service.document_repo.create.assert_not_called()
    service.index.delete_by_document_id.assert_called_once_with(
        "failed-document-id"
    )
    service.document_repo.mark_processing.assert_called_once_with(
        "failed-document-id"
    )
    service.extract.run.assert_called_once_with(
        source_data=pdf_data,
        filename="lesson.pdf",
        file_type="pdf",
        document_id="failed-document-id",
        subject_id="subject-id",
    )
    service.transform.run.assert_called_once_with(
        extracted
    )
    service.chunk.run.assert_called_once_with(
        transformed
    )
    service.embedding.run.assert_called_once_with(
        chunks
    )
    service.index.run.assert_called_once_with(
        chunks,
        embeddings,
    )
    service.document_repo.mark_indexed.assert_called_once_with(
        "failed-document-id"
    )


def test_retry_rejects_source_with_changed_content(
    tmp_path,
):
    original_data = b"%PDF-1.4\noriginal"
    changed_data = b"%PDF-1.4\ntampered"

    assert len(original_data) == len(changed_data)

    source_path = tmp_path / "retained.pdf"
    source_path.write_bytes(changed_data)

    document = SimpleNamespace(
        document_id="failed-document-id",
        subject_id="subject-id",
        filename="lesson.pdf",
        status="FAILED",
        source_uri=source_path.as_posix(),
        content_sha256=sha256(
            original_data
        ).hexdigest(),
        file_size_bytes=len(original_data),
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = LocalSourceStorage(
        root=tmp_path,
    )
    service.document_repo = Mock()
    service.index = Mock()

    with pytest.raises(
        RuntimeError,
        match="checksum does not match",
    ):
        service.retry_failed_document(document)

    service.document_repo.mark_processing.assert_not_called()
    service.index.delete_by_document_id.assert_not_called()


def test_processing_failure_records_failed_stage():
    service = PDFService.__new__(PDFService)

    service.document_repo = Mock()
    service.extract = Mock()
    service.transform = Mock()
    service.chunk = Mock()
    service.embedding = Mock()
    service.index = Mock()

    extracted = [object()]

    service.extract.run.return_value = extracted
    service.transform.run.side_effect = RuntimeError(
        "Controlled transform failure"
    )

    with pytest.raises(
        RuntimeError,
        match="Controlled transform failure",
    ):
        service._process_document(
            source_data=b"%PDF-1.4\nretained",
            filename="retained.pdf",
            file_type="pdf",
            document_id="failed-document-id",
            subject_id="subject-id",
        )

    service.document_repo.mark_processing.assert_called_once_with(
        "failed-document-id"
    )
    service.document_repo.mark_failed.assert_called_once_with(
        document_id="failed-document-id",
        failed_stage="TRANSFORM",
        error_message="Controlled transform failure",
    )
    service.document_repo.mark_indexed.assert_not_called()
    service.chunk.run.assert_not_called()


def test_index_failure_cleans_up_vectors_before_marking_failed():
    service = PDFService.__new__(PDFService)

    service.document_repo = Mock()
    service.extract = Mock()
    service.transform = Mock()
    service.chunk = Mock()
    service.embedding = Mock()
    service.index = Mock()

    service.extract.run.return_value = [object()]
    service.transform.run.return_value = [object()]
    service.chunk.run.return_value = [object()]
    service.embedding.run.return_value = object()

    call_log = []

    def fail_index(*_args):
        call_log.append("index_run")
        raise RuntimeError("Controlled index failure")

    def cleanup_vectors(document_id):
        assert document_id == "failed-document-id"
        call_log.append("vector_cleanup")

    def mark_failed(**_kwargs):
        call_log.append("mark_failed")

    service.index.run.side_effect = fail_index
    service.index.delete_by_document_id.side_effect = (
        cleanup_vectors
    )
    service.document_repo.mark_failed.side_effect = mark_failed

    with pytest.raises(
        RuntimeError,
        match="Controlled index failure",
    ):
        service._process_document(
            source_data=b"%PDF-1.4\nretained",
            filename="retained.pdf",
            file_type="pdf",
            document_id="failed-document-id",
            subject_id="subject-id",
        )

    assert call_log == [
        "index_run",
        "vector_cleanup",
        "mark_failed",
    ]
    service.index.delete_by_document_id.assert_called_once_with(
        "failed-document-id"
    )
    service.document_repo.mark_failed.assert_called_once_with(
        document_id="failed-document-id",
        failed_stage="INDEX",
        error_message="Controlled index failure",
    )


def test_index_failure_marks_failed_when_vector_cleanup_fails():
    service = PDFService.__new__(PDFService)

    service.document_repo = Mock()
    service.extract = Mock()
    service.transform = Mock()
    service.chunk = Mock()
    service.embedding = Mock()
    service.index = Mock()

    service.extract.run.return_value = [object()]
    service.transform.run.return_value = [object()]
    service.chunk.run.return_value = [object()]
    service.embedding.run.return_value = object()
    service.index.run.side_effect = RuntimeError(
        "Controlled index failure"
    )
    service.index.delete_by_document_id.side_effect = (
        RuntimeError("Controlled cleanup failure")
    )

    with pytest.raises(
        RuntimeError,
        match="Controlled index failure",
    ):
        service._process_document(
            source_data=b"%PDF-1.4\nretained",
            filename="retained.pdf",
            file_type="pdf",
            document_id="failed-document-id",
            subject_id="subject-id",
        )

    service.index.delete_by_document_id.assert_called_once_with(
        "failed-document-id"
    )
    service.document_repo.mark_failed.assert_called_once_with(
        document_id="failed-document-id",
        failed_stage="INDEX",
        error_message=(
            "Controlled index failure; derived vector cleanup "
            "failed: Controlled cleanup failure"
        ),
    )


@pytest.mark.parametrize(
    "status",
    [
        "UPLOADED",
        "PROCESSING",
    ],
)
def test_recover_stale_document_reuses_source_and_document_id(
    tmp_path,
    status,
):
    pdf_data = b"%PDF-1.4\nstale retained source"
    source_path = tmp_path / "stale.pdf"
    source_path.write_bytes(pdf_data)

    document = SimpleNamespace(
        document_id="stale-document-id",
        subject_id="subject-id",
        file_type="pdf",
        filename="lesson.pdf",
        status=status,
        source_uri=source_path.as_posix(),
        content_sha256=sha256(pdf_data).hexdigest(),
        file_size_bytes=len(pdf_data),
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = LocalSourceStorage(
        root=tmp_path,
    )
    service.document_repo = Mock()
    service.index = Mock()
    service._process_document = Mock(
        return_value=2
    )

    result = service.recover_stale_document(
        document
    )

    assert result == 2

    service.document_repo.create.assert_not_called()
    service.index.delete_by_document_id.assert_called_once_with(
        "stale-document-id"
    )
    service._process_document.assert_called_once_with(
        source_data=pdf_data,
        filename="lesson.pdf",
        file_type="pdf",
        document_id="stale-document-id",
        subject_id="subject-id",
    )


@pytest.mark.parametrize(
    "status",
    [
        "FAILED",
        "INDEXED",
        "DELETED",
    ],
)
def test_recover_stale_document_rejects_other_statuses(
    status,
):
    document = SimpleNamespace(
        status=status,
    )

    service = PDFService.__new__(PDFService)

    with pytest.raises(
        RuntimeError,
        match=(
            "Only stale uploaded or processing "
            "documents can be recovered"
        ),
    ):
        service.recover_stale_document(
            document
        )


def test_recover_stale_document_marks_missing_source_failed():
    document = SimpleNamespace(
        document_id="legacy-document-id",
        subject_id="subject-id",
        status="UPLOADED",
        file_type="pdf",
        source_uri=None,
        content_sha256=None,
        file_size_bytes=None,
    )

    service = PDFService.__new__(PDFService)
    service.document_repo = Mock()
    service.index = Mock()

    with pytest.raises(
        RuntimeError,
        match="Document source metadata is incomplete",
    ):
        service.recover_stale_document(
            document
        )

    service.document_repo.mark_failed.assert_called_once_with(
        document_id="legacy-document-id",
        failed_stage="EXTRACT",
        error_message=(
            "Document source metadata is incomplete"
        ),
    )
    service.index.delete_by_document_id.assert_not_called()


@pytest.mark.parametrize(
    "status",
    [
        "UPLOADED",
        "PROCESSING",
        "FAILED",
        "INDEXED",
    ],
)
def test_rebuild_retained_document_accepts_non_deleted_statuses(
    status,
):
    document = SimpleNamespace(
        document_id="document-id",
        status=status,
    )

    service = PDFService.__new__(
        PDFService
    )
    service._reprocess_retained_document = Mock(
        return_value=3
    )

    result = service.rebuild_retained_document(
        document
    )

    assert result == 3
    (
        service
        ._reprocess_retained_document
        .assert_called_once_with(document)
    )


def test_rebuild_retained_document_rejects_deleted_document():
    document = SimpleNamespace(
        document_id="deleted-document-id",
        status="DELETED",
    )

    service = PDFService.__new__(
        PDFService
    )
    service._reprocess_retained_document = Mock()

    with pytest.raises(
        RuntimeError,
        match=(
            "Document cannot be rebuilt "
            "from its current status"
        ),
    ):
        service.rebuild_retained_document(
            document
        )

    (
        service
        ._reprocess_retained_document
        .assert_not_called()
    )


def test_rebuild_retained_document_marks_invalid_source_failed():
    document = SimpleNamespace(
        document_id="indexed-document-id",
        subject_id="subject-id",
        status="INDEXED",
        source_uri=None,
        content_sha256=None,
        file_size_bytes=None,
    )

    service = PDFService.__new__(
        PDFService
    )
    service.document_repo = Mock()
    service.index = Mock()

    with pytest.raises(
        RuntimeError,
        match=(
            "Document source metadata "
            "is incomplete"
        ),
    ):
        service.rebuild_retained_document(
            document
        )

    service.document_repo.mark_failed.assert_called_once_with(
        document_id="indexed-document-id",
        failed_stage="EXTRACT",
        error_message=(
            "Document source metadata "
            "is incomplete"
        ),
    )
    service.index.delete_by_document_id.assert_not_called()


def test_upload_reserves_document_before_storing_source():
    text_data = b"first line\nsecond line"
    stored_source = StoredSource(
        source_uri="source://document-1",
        content_sha256=sha256(text_data).hexdigest(),
        file_size_bytes=len(text_data),
    )

    source_storage = Mock()
    call_order = []

    def reserve_document(**_kwargs):
        call_order.append("reserve")

    def store_source(**_kwargs):
        call_order.append("store")
        return stored_source

    def attach_metadata(**_kwargs):
        call_order.append("attach")

    def process_document(**_kwargs):
        call_order.append("process")
        return 2

    service = PDFService.__new__(PDFService)
    service.source_storage = source_storage
    service.document_repo = Mock()
    service.document_repo.create.side_effect = (
        reserve_document
    )
    service.document_repo.attach_source_metadata.side_effect = (
        attach_metadata
    )
    service.source_storage.store.side_effect = store_source
    service._process_document = Mock(
        side_effect=process_document
    )

    result = service.upload(
        source_data=text_data,
        subject_id="subject-1",
        document_id="document-1",
        original_filename="lesson.txt",
        content_sha256=sha256(text_data).hexdigest(),
        file_size_bytes=len(text_data),
    )

    assert result == 2
    assert call_order == [
        "reserve",
        "store",
        "attach",
        "process",
    ]

    service.document_repo.create.assert_called_once_with(
        document_id="document-1",
        subject_id="subject-1",
        filename="lesson.txt",
        file_type="txt",
        chapter=None,
        source_uri=None,
        content_sha256=sha256(text_data).hexdigest(),
        file_size_bytes=len(text_data),
    )
    service.source_storage.store.assert_called_once_with(
        document_id="document-1",
        filename="lesson.txt",
        data=text_data,
    )
    service.document_repo.attach_source_metadata.assert_called_once_with(
        document_id="document-1",
        source_uri="source://document-1",
        content_sha256=sha256(text_data).hexdigest(),
        file_size_bytes=len(text_data),
    )
    source_storage.read_bytes.assert_not_called()


def test_upload_does_not_store_source_when_reservation_fails():
    text_data = b"first line\nsecond line"

    service = PDFService.__new__(PDFService)
    service.source_storage = Mock()
    service.document_repo = Mock()
    service.document_repo.create.side_effect = RuntimeError(
        "document reservation conflict"
    )
    service._process_document = Mock()

    with pytest.raises(
        RuntimeError,
        match="document reservation conflict",
    ):
        service.upload(
            source_data=text_data,
            subject_id="subject-1",
            document_id="document-1",
            original_filename="lesson.txt",
            content_sha256=sha256(text_data).hexdigest(),
            file_size_bytes=len(text_data),
        )

    service.source_storage.store.assert_not_called()
    service.document_repo.attach_source_metadata.assert_not_called()
    service._process_document.assert_not_called()


def test_process_txt_preserves_exact_line_ranges(
    tmp_path,
):
    source_path = tmp_path / "lesson.txt"
    source_path.write_bytes(
        b"alpha\r\nbeta\r\ngamma\r\n"
    )

    service = PDFService.__new__(PDFService)
    service.document_repo = Mock()
    service.extract = ExtractPipeline()
    service.transform = TransformPipeline()
    service.chunk = ChunkPipeline()
    service.chunk.chunker = TextChunker(
        chunk_size=6,
        overlap=0,
    )
    service.embedding = Mock()
    service.index = Mock()

    embeddings = object()
    service.embedding.run.return_value = embeddings

    result = service._process_document(
        source_data=source_path.read_bytes(),
        filename="lesson.txt",
        file_type="txt",
        document_id="document-1",
        subject_id="subject-1",
    )

    assert result == 3

    chunks = service.index.run.call_args.args[0]

    assert [chunk.text for chunk in chunks] == [
        "alpha\n",
        "beta\ng",
        "amma",
    ]
    assert [
        chunk.line_range for chunk in chunks
    ] == [
        (1, 1),
        (2, 3),
        (3, 3),
    ]
    assert all(
        chunk.page is None
        for chunk in chunks
    )
    assert all(
        chunk.paragraph_range is None
        for chunk in chunks
    )

    service.embedding.run.assert_called_once_with(
        chunks
    )
    service.index.run.assert_called_once_with(
        chunks,
        embeddings,
    )
    service.document_repo.mark_indexed.assert_called_once_with(
        "document-1"
    )
    service.document_repo.mark_failed.assert_not_called()


def test_process_txt_reuses_chunk_ids_for_same_document():
    source_data = b"alpha\nbeta\ngamma\n"

    service = PDFService.__new__(PDFService)
    service.document_repo = Mock()
    service.extract = ExtractPipeline()
    service.transform = TransformPipeline()
    service.chunk = ChunkPipeline()
    service.chunk.chunker = TextChunker(
        chunk_size=6,
        overlap=0,
    )
    service.embedding = Mock()
    service.index = Mock()
    service.embedding.run.return_value = object()

    for _ in range(2):
        service._process_document(
            source_data=source_data,
            filename="lesson.txt",
            file_type="txt",
            document_id="document-1",
            subject_id="subject-1",
        )

    first_chunks = service.index.run.call_args_list[0].args[0]
    second_chunks = service.index.run.call_args_list[1].args[0]

    assert [chunk.id for chunk in first_chunks] == [
        chunk.id for chunk in second_chunks
    ]


def test_rebuild_retained_txt_reuses_identity_and_lineage(
    tmp_path,
):
    text_data = b"alpha\nbeta\ngamma"
    source_path = tmp_path / "retained.txt"
    source_path.write_bytes(text_data)

    document = SimpleNamespace(
        document_id="document-1",
        subject_id="subject-1",
        filename="lesson.txt",
        file_type="txt",
        status="INDEXED",
        source_uri=source_path.as_posix(),
        content_sha256=sha256(
            text_data
        ).hexdigest(),
        file_size_bytes=len(text_data),
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = LocalSourceStorage(
        root=tmp_path,
    )
    service.document_repo = Mock()
    service.extract = ExtractPipeline()
    service.transform = TransformPipeline()
    service.chunk = ChunkPipeline()
    service.chunk.chunker = TextChunker(
        chunk_size=6,
        overlap=0,
    )
    service.embedding = Mock()
    service.index = Mock()

    embeddings = object()
    service.embedding.run.return_value = embeddings

    result = service.rebuild_retained_document(
        document
    )

    assert result == 3

    service.index.delete_by_document_id.assert_called_once_with(
        "document-1"
    )
    service.document_repo.create.assert_not_called()

    chunks = service.index.run.call_args.args[0]

    assert {
        chunk.document_id
        for chunk in chunks
    } == {"document-1"}
    assert [
        chunk.line_range
        for chunk in chunks
    ] == [
        (1, 1),
        (2, 3),
        (3, 3),
    ]

    service.document_repo.mark_processing.assert_called_once_with(
        "document-1"
    )
    service.document_repo.mark_indexed.assert_called_once_with(
        "document-1"
    )


def test_process_docx_preserves_exact_paragraph_ranges(
    tmp_path,
):
    source_path = tmp_path / "lesson.docx"

    word_document = WordDocument()
    word_document.add_paragraph("alpha")
    word_document.add_paragraph("beta")
    word_document.add_paragraph("gamma")
    word_document.save(source_path)

    service = PDFService.__new__(PDFService)
    service.document_repo = Mock()
    service.extract = ExtractPipeline()
    service.transform = TransformPipeline()
    service.chunk = ChunkPipeline()
    service.chunk.chunker = TextChunker(
        chunk_size=6,
        overlap=0,
    )
    service.embedding = Mock()
    service.index = Mock()

    embeddings = object()
    service.embedding.run.return_value = embeddings

    result = service._process_document(
        source_data=source_path.read_bytes(),
        filename="lesson.docx",
        file_type="docx",
        document_id="document-1",
        subject_id="subject-1",
    )

    assert result == 3

    chunks = service.index.run.call_args.args[0]

    assert [chunk.text for chunk in chunks] == [
        "alpha\n",
        "beta\ng",
        "amma",
    ]
    assert [
        chunk.paragraph_range
        for chunk in chunks
    ] == [
        (1, 1),
        (2, 3),
        (3, 3),
    ]
    assert all(
        chunk.page is None
        for chunk in chunks
    )
    assert all(
        chunk.line_range is None
        for chunk in chunks
    )

    service.document_repo.mark_indexed.assert_called_once_with(
        "document-1"
    )
    service.document_repo.mark_failed.assert_not_called()


def test_rebuild_retained_docx_reuses_identity_and_lineage(
    tmp_path,
):
    source_path = tmp_path / "retained.docx"

    word_document = WordDocument()
    word_document.add_paragraph("alpha")
    word_document.add_paragraph("beta")
    word_document.add_paragraph("gamma")
    word_document.save(source_path)

    source_data = source_path.read_bytes()

    document = SimpleNamespace(
        document_id="document-1",
        subject_id="subject-1",
        filename="lesson.docx",
        file_type="docx",
        status="INDEXED",
        source_uri=source_path.as_posix(),
        content_sha256=sha256(
            source_data
        ).hexdigest(),
        file_size_bytes=len(source_data),
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = LocalSourceStorage(
        root=tmp_path,
    )
    service.document_repo = Mock()
    service.extract = ExtractPipeline()
    service.transform = TransformPipeline()
    service.chunk = ChunkPipeline()
    service.chunk.chunker = TextChunker(
        chunk_size=6,
        overlap=0,
    )
    service.embedding = Mock()
    service.index = Mock()

    embeddings = object()
    service.embedding.run.return_value = embeddings

    result = service.rebuild_retained_document(
        document
    )

    assert result == 3

    service.index.delete_by_document_id.assert_called_once_with(
        "document-1"
    )
    service.document_repo.create.assert_not_called()

    chunks = service.index.run.call_args.args[0]

    assert {
        chunk.document_id
        for chunk in chunks
    } == {"document-1"}
    assert [
        chunk.paragraph_range
        for chunk in chunks
    ] == [
        (1, 1),
        (2, 3),
        (3, 3),
    ]
    assert all(
        chunk.line_range is None
        for chunk in chunks
    )
    assert all(
        chunk.page is None
        for chunk in chunks
    )

    service.document_repo.mark_processing.assert_called_once_with(
        "document-1"
    )
    service.document_repo.mark_indexed.assert_called_once_with(
        "document-1"
    )


def test_rebuild_reads_retained_source_through_storage():
    source_data = b"retained source"

    document = SimpleNamespace(
        document_id="document-1",
        subject_id="subject-1",
        filename="lesson.txt",
        file_type="txt",
        status="INDEXED",
        source_uri="source://document-1",
        content_sha256=(
            sha256(source_data).hexdigest()
        ),
        file_size_bytes=len(source_data),
    )

    source_storage = Mock()
    source_storage.read_bytes.return_value = (
        source_data
    )

    service = PDFService.__new__(PDFService)
    service.document_repo = Mock()
    service.index = Mock()
    service.source_storage = source_storage
    service._process_document = Mock(
        return_value=3
    )

    result = service.rebuild_retained_document(
        document
    )

    assert result == 3
    source_storage.read_bytes.assert_called_once_with(
        "source://document-1"
    )
    service._process_document.assert_called_once_with(
        source_data=source_data,
        filename="lesson.txt",
        file_type="txt",
        document_id="document-1",
        subject_id="subject-1",
    )


def test_pdf_service_accepts_injected_source_storage():
    source_storage = Mock()

    bound_arguments = signature(
        PDFService
    ).bind(
        Mock(),
        source_storage=source_storage,
    )

    assert (
        bound_arguments.arguments[
            "source_storage"
        ]
        is source_storage
    )


def test_rebuild_marks_missing_storage_source_failed_before_deleting_index():
    document = SimpleNamespace(
        document_id="document-1",
        subject_id="subject-1",
        filename="lesson.txt",
        file_type="txt",
        status="INDEXED",
        source_uri=(
            "minio://cloudmentor/"
            "documents/document-1/source.txt"
        ),
        content_sha256=sha256(
            b"expected source"
        ).hexdigest(),
        file_size_bytes=len(
            b"expected source"
        ),
    )

    source_storage = Mock()
    source_storage.read_bytes.side_effect = (
        SourceNotFoundError(
            "Stored source does not exist"
        )
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = source_storage
    service.document_repo = Mock()
    service.index = Mock()

    with pytest.raises(
        RetainedSourceValidationError,
        match="Retained source file is missing",
    ):
        service.rebuild_retained_document(
            document
        )

    service.document_repo.mark_failed.assert_called_once_with(
        document_id="document-1",
        failed_stage="EXTRACT",
        error_message=(
            "Retained source file is missing"
        ),
    )
    service.document_repo.mark_processing.assert_not_called()
    service.index.delete_by_document_id.assert_not_called()


def test_rebuild_preserves_document_and_index_when_storage_unavailable():
    document = SimpleNamespace(
        document_id="document-1",
        subject_id="subject-1",
        filename="lesson.txt",
        file_type="txt",
        status="INDEXED",
        source_uri=(
            "minio://cloudmentor/"
            "documents/document-1/source.txt"
        ),
        content_sha256=sha256(
            b"expected source"
        ).hexdigest(),
        file_size_bytes=len(
            b"expected source"
        ),
    )

    source_storage = Mock()
    source_storage.read_bytes.side_effect = (
        SourceUnavailableError(
            "Source storage is temporarily unavailable"
        )
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = source_storage
    service.document_repo = Mock()
    service.index = Mock()

    with pytest.raises(
        SourceUnavailableError,
        match="temporarily unavailable",
    ):
        service.rebuild_retained_document(
            document
        )

    service.document_repo.mark_failed.assert_not_called()
    service.document_repo.mark_processing.assert_not_called()
    service.index.delete_by_document_id.assert_not_called()


def test_rebuild_rejects_checksum_mismatch_before_deleting_index():
    expected_data = b"expected source"
    changed_data = b"tampered source"

    assert len(expected_data) == len(
        changed_data
    )

    document = SimpleNamespace(
        document_id="document-1",
        subject_id="subject-1",
        filename="lesson.txt",
        file_type="txt",
        status="INDEXED",
        source_uri=(
            "minio://cloudmentor/"
            "documents/document-1/source.txt"
        ),
        content_sha256=sha256(
            expected_data
        ).hexdigest(),
        file_size_bytes=len(expected_data),
    )

    source_storage = Mock()
    source_storage.read_bytes.return_value = (
        changed_data
    )

    service = PDFService.__new__(PDFService)
    service.source_storage = source_storage
    service.document_repo = Mock()
    service.index = Mock()

    with pytest.raises(
        RetainedSourceValidationError,
        match="checksum does not match",
    ):
        service.rebuild_retained_document(
            document
        )

    service.document_repo.mark_failed.assert_called_once_with(
        document_id="document-1",
        failed_stage="EXTRACT",
        error_message=(
            "Retained source checksum does not match"
        ),
    )
    service.document_repo.mark_processing.assert_not_called()
    service.index.delete_by_document_id.assert_not_called()
