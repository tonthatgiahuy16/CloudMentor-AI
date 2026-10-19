from io import BytesIO
from typing import List

from pypdf import PdfReader

from app.core.logger import logger
from app.models.document import Document
from docx import Document as WordDocument
from zipfile import BadZipFile

from docx.opc.exceptions import (
    PackageNotFoundError,
)

def normalize_docx_paragraph_text(
    text: str,
) -> str:
    return (
        text
        .replace("\r\n", " ")
        .replace("\r", " ")
        .replace("\n", " ")
    )


def read_docx_paragraphs(source) -> list[str]:
    try:
        word_document = WordDocument(source)
    except (
        BadZipFile,
        PackageNotFoundError,
    ) as exc:
        raise ValueError(
            "DOCX file is not valid"
        ) from exc

    if any(
        cell.text.strip()
        for table in word_document.tables
        for row in table.rows
        for cell in row.cells
    ):
        raise ValueError(
            "DOCX tables are not supported"
        )

    paragraphs = [
        normalize_docx_paragraph_text(
        paragraph.text
        )
        for paragraph in word_document.paragraphs
    ]

    if not paragraphs or not any(
        paragraph.strip()
        for paragraph in paragraphs
    ):
        raise ValueError(
            "DOCX file is empty"
        )

    return paragraphs

class PDFLoader:
    """Read PDF bytes into page-level documents."""

    def load(
        self,
        source_data: bytes,
        *,
        filename: str,
        document_id: str,
        subject_id: str,
    ) -> List[Document]:
        if not filename.lower().endswith(".pdf"):
            raise ValueError(
                f"Expected a PDF file: {filename}"
            )

        logger.info(
            "Loading PDF: %s",
            filename,
        )

        reader = PdfReader(
            BytesIO(source_data)
        )

        documents = [
            Document(
                document_id=document_id,
                subject_id=subject_id,
                page=page_number,
                source=filename,
                text=page.extract_text() or "",
            )
            for page_number, page in enumerate(
                reader.pages,
                start=1,
            )
        ]

        logger.info(
            "Loaded %s pages from %s",
            len(documents),
            filename,
        )

        return documents


class TextLoader:
    """Read UTF-8 TXT bytes with line-level lineage."""

    def load(
        self,
        source_data: bytes,
        *,
        filename: str,
        document_id: str,
        subject_id: str,
    ) -> List[Document]:
        if not filename.lower().endswith(".txt"):
            raise ValueError(
                f"Expected a TXT file: {filename}"
            )

        logger.info(
            "Loading TXT: %s",
            filename,
        )

        try:
            text = source_data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ValueError(
                "TXT file is not valid UTF-8: "
                f"{filename}"
            ) from exc

        lines = text.splitlines()

        if not lines or not any(
            line.strip()
            for line in lines
        ):
            raise ValueError(
                f"TXT file is empty: {filename}"
            )

        normalized_text = "\n".join(lines)

        document = Document(
            document_id=document_id,
            subject_id=subject_id,
            source=filename,
            text=normalized_text,
            line_range=(1, len(lines)),
        )

        logger.info(
            "Loaded %s lines from %s",
            len(lines),
            filename,
        )

        return [document]


class DOCXLoader:
    """Read DOCX bytes with paragraph-level lineage."""

    def load(
        self,
        source_data: bytes,
        *,
        filename: str,
        document_id: str,
        subject_id: str,
    ) -> List[Document]:
        if not filename.lower().endswith(".docx"):
            raise ValueError(
                f"Expected a DOCX file: {filename}"
            )

        logger.info(
            "Loading DOCX: %s",
            filename,
        )

        paragraphs = read_docx_paragraphs(
            BytesIO(source_data)
        )

        document = Document(
            document_id=document_id,
            subject_id=subject_id,
            source=filename,
            text="\n".join(paragraphs),
            paragraph_range=(
                1,
                len(paragraphs),
            ),
        )

        logger.info(
            "Loaded %s paragraphs from %s",
            len(paragraphs),
            filename,
        )

        return [document]
